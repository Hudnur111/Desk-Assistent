import asyncio
from pathlib import Path

from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation

from .base import Tool

DOCUMENTS_DIR = Path("data/documents")


def _safe_path(title: str, suffix: str = ".docx") -> Path:
    DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)
    return DOCUMENTS_DIR / f"{title}{suffix}"


async def _create_document(title: str, content: str) -> str:
    def _write() -> Path:
        path = _safe_path(title)
        document = Document()
        document.add_heading(title, level=1)
        for paragraph in content.split("\n\n"):
            document.add_paragraph(paragraph)
        document.save(path)
        return path

    path = await asyncio.to_thread(_write)
    return f"Dokument erstellt: {path}"


async def _append_to_document(path: str, text: str) -> str:
    file_path = Path(path)
    if not file_path.is_file():
        return f"Dokument nicht gefunden: {path}"

    def _write() -> None:
        document = Document(file_path)
        document.add_paragraph(text)
        document.save(file_path)

    await asyncio.to_thread(_write)
    return f"Text an {path} angehaengt"


async def _read_document(path: str) -> str:
    file_path = Path(path)

    def _read() -> str | None:
        if not file_path.is_file():
            return None
        document = Document(file_path)
        return "\n".join(p.text for p in document.paragraphs)

    content = await asyncio.to_thread(_read)
    if content is None:
        return f"Dokument nicht gefunden: {path}"
    return content


async def _create_spreadsheet(title: str, rows: list[list[str]]) -> str:
    def _write() -> Path:
        path = _safe_path(title, ".xlsx")
        wb = Workbook()
        ws = wb.active
        for row in rows:
            ws.append(row)
        wb.save(path)
        return path

    path = await asyncio.to_thread(_write)
    return f"Excel-Tabelle erstellt: {path}"


async def _read_spreadsheet(path: str) -> str:
    file_path = Path(path)

    def _read() -> str | None:
        if not file_path.is_file():
            return None
        wb = load_workbook(file_path)
        ws = wb.active
        return "\n".join(
            ", ".join("" if cell is None else str(cell) for cell in row)
            for row in ws.iter_rows(values_only=True)
        )

    content = await asyncio.to_thread(_read)
    if content is None:
        return f"Tabelle nicht gefunden: {path}"
    return content


async def _create_presentation(title: str, slides: list[dict[str, str]]) -> str:
    def _write() -> Path:
        path = _safe_path(title, ".pptx")
        prs = Presentation()
        title_layout = prs.slide_layouts[0]
        content_layout = prs.slide_layouts[1]

        first = prs.slides.add_slide(title_layout)
        first.shapes.title.text = title

        for slide_data in slides:
            slide = prs.slides.add_slide(content_layout)
            slide.shapes.title.text = slide_data.get("title", "")
            body = slide.placeholders[1].text_frame
            body.text = slide_data.get("content", "")
        prs.save(path)
        return path

    path = await asyncio.to_thread(_write)
    return f"PowerPoint-Praesentation erstellt: {path}"


def document_tools() -> list[Tool]:
    return [
        Tool(
            name="create_document",
            description="Erstellt ein neues Word-Dokument (.docx) mit Titel und Inhalt.",
            input_schema={
                "type": "object",
                "properties": {
                    "title": {
                        "type": "string",
                        "description": "Titel des Dokuments, wird auch als Dateiname verwendet",
                    },
                    "content": {
                        "type": "string",
                        "description": "Inhalt; durch Leerzeilen getrennte Abschnitte werden als eigene Absaetze angelegt",
                    },
                },
                "required": ["title", "content"],
            },
            handler=_create_document,
        ),
        Tool(
            name="append_to_document",
            description="Haengt einen weiteren Absatz an ein bestehendes Word-Dokument an.",
            input_schema={
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "Pfad zur .docx-Datei"},
                    "text": {"type": "string", "description": "Anzuhaengender Absatz"},
                },
                "required": ["path", "text"],
            },
            handler=_append_to_document,
        ),
        Tool(
            name="read_document",
            description="Liest den Textinhalt eines Word-Dokuments (.docx).",
            input_schema={
                "type": "object",
                "properties": {
                    "path": {"type": "string", "description": "Pfad zur .docx-Datei"}
                },
                "required": ["path"],
            },
            handler=_read_document,
        ),
        Tool(
            name="create_spreadsheet",
            description="Erstellt eine neue Excel-Tabelle (.xlsx) aus Zeilen von Zellwerten.",
            input_schema={
                "type": "object",
                "properties": {
                    "title": {"type": "string", "description": "Titel/Dateiname der Tabelle"},
                    "rows": {
                        "type": "array",
                        "items": {"type": "array", "items": {"type": "string"}},
                        "description": "Liste von Zeilen, jede Zeile eine Liste von Zellwerten",
                    },
                },
                "required": ["title", "rows"],
            },
            handler=_create_spreadsheet,
        ),
        Tool(
            name="read_spreadsheet",
            description="Liest den Inhalt einer Excel-Tabelle (.xlsx) als Text.",
            input_schema={
                "type": "object",
                "properties": {"path": {"type": "string", "description": "Pfad zur .xlsx-Datei"}},
                "required": ["path"],
            },
            handler=_read_spreadsheet,
        ),
        Tool(
            name="create_presentation",
            description="Erstellt eine neue PowerPoint-Praesentation (.pptx) mit Titelfolie und Inhaltsfolien.",
            input_schema={
                "type": "object",
                "properties": {
                    "title": {"type": "string", "description": "Titel der Praesentation"},
                    "slides": {
                        "type": "array",
                        "items": {
                            "type": "object",
                            "properties": {
                                "title": {"type": "string"},
                                "content": {"type": "string"},
                            },
                        },
                        "description": "Liste von Folien mit title und content",
                    },
                },
                "required": ["title", "slides"],
            },
            handler=_create_presentation,
        ),
    ]
